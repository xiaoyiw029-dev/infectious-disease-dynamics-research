from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SRC = Path(r"D:\传染病动力学科研\06_论文与报告\04_中期答辩\中期答辩_修改完善版_v03.docx")
OUT = Path(r"D:\传染病动力学科研\06_论文与报告\04_中期答辩\中期答辩_修改完善版_v04.docx")


TARGET_PICTURE_ORDER = {
    2: ["[[EQ_FORCE]]"],
    3: [
        "[[EQ_M1]]",
        "[[EQ_M2]]",
        "[[EQ_M3]]",
        "[[EQ_M4]]",
        "[[EQ_M5]]",
        "[[EQ_M6]]",
        "[[EQ_M7]]",
    ],
    4: ["[[EQ_DFE]]"],
    5: ["[[EQ_R0]]"],
    7: ["[[EQ_OBS]]"],
}


def has_inline_picture(paragraph):
    return bool(paragraph._p.xpath(".//wp:inline"))


def replace_paragraph_text(paragraph, text):
    pPr = paragraph._p.pPr
    preserved = deepcopy(pPr) if pPr is not None else None
    paragraph._p.clear_content()
    if preserved is not None:
        paragraph._p.insert(0, preserved)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)


def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted.paragraph_format.space_before = Pt(3)
    inserted.paragraph_format.space_after = Pt(3)
    run = inserted.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)
    return inserted


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    doc = Document(SRC)
    cell = doc.tables[0].rows[11].cells[1]
    pic_order = 0
    replacements = 0

    for paragraph in list(cell.paragraphs):
        if not has_inline_picture(paragraph):
            continue

        pic_order += 1
        placeholders = TARGET_PICTURE_ORDER.get(pic_order)
        if not placeholders:
            continue

        replace_paragraph_text(paragraph, placeholders[0])
        cursor = paragraph
        for placeholder in placeholders[1:]:
            cursor = insert_after(cursor, placeholder)
        replacements += 1

    missing = set(TARGET_PICTURE_ORDER) - set(range(1, pic_order + 1))
    if missing or replacements != len(TARGET_PICTURE_ORDER):
        raise RuntimeError(
            f"Formula image replacement mismatch: pictures={pic_order}, "
            f"replacements={replacements}, missing={sorted(missing)}"
        )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
