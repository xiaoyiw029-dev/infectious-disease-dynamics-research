from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SRC = Path(r"C:\Users\Lenovo\OneDrive\Desktop\巴西基孔肯亚热创新训练计划项目进展报告书.docx")
OUT = Path(r"D:\传染病动力学科研\06_论文与报告\04_中期答辩\巴西基孔肯亚热创新训练计划项目进展报告书_格式校正版_v01.docx")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(width_twips))
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "center")


def disable_table_autofit(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_row_keep_together(row, keep=True):
    trPr = row._tr.get_or_add_trPr()
    cant_split = trPr.find(qn("w:cantSplit"))
    if keep and cant_split is None:
        trPr.append(OxmlElement("w:cantSplit"))
    elif not keep and cant_split is not None:
        trPr.remove(cant_split)


def set_paragraph_basics(paragraph):
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = run.font.name or "宋体"


def fit_inline_images(doc, max_width_cm):
    max_width = Cm(max_width_cm)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    doc = Document(SRC)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    section.start_type = WD_SECTION.NEW_PAGE

    usable_width_cm = 21 - 1.4 - 1.4
    usable_width_twips = int(usable_width_cm / 2.54 * 1440)

    for idx, paragraph in enumerate(doc.paragraphs):
        set_paragraph_basics(paragraph)
        if idx == 0 and paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(16)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        disable_table_autofit(table)
        set_table_width(table, usable_width_twips)

        for row_index, row in enumerate(table.rows):
            # Keep compact heading/signature rows together. Allow the long body rows to split.
            set_row_keep_together(row, keep=row_index not in {11, 12})
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                for p in cell.paragraphs:
                    set_paragraph_basics(p)
                    text = p.text.strip()
                    if text.startswith("图 ") or text.startswith("表 "):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if text in {"项目进展：", "主要参考文献：", "存在的主要研究难题：", "后续研究计划："}:
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(10)

    fit_inline_images(doc, max_width_cm=14.8)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
